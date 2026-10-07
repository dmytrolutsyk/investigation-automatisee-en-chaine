"""Rapport d'incident Word (python-docx, A4), dans la charte Formind.

Ce module ne fait que la mise en page : les textes arrivent déjà rédigés dans
le dictionnaire produit par investigation.construire_donnees_rapport. Même
contenu organisé en sections, avec une page de garde.
"""
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from affichage import (DECOMPTE_IOC as _DECOMPTE_IOC, EN_TETE_PRIORITE as _EN_TETE_PRIORITE,
                       LIBELLES as _LIBELLES, PRUDENCE, tronquer)

# --- Charte -----------------------------------------------------------------
MARINE, TURQUOISE = "002236", "0089A4"   # relevés sur la charte Formind
LIGNE_CLAIRE, FOND_ALERTE, ALERTE = "F2F7F8", "FBEAE8", "C0392B"
TEXTE, GRIS = "22303C", "6B7A86"
COULEURS_GRAVITE = {"CRITIQUE": "C0392B", "ÉLEVÉE": "E67E22", "MODÉRÉE": "F1C40F",
                    "FAIBLE": "27AE60"}
COULEURS_PRIORITE = {"IMMÉDIAT": "C0392B", "COURT TERME": "E67E22", "SUIVI": TURQUOISE}
POLICE = "Poppins"          # police de la charte Formind
POLICE_REPLI = "Arial"      # proposée par Word si Poppins n'est pas installée
LARGEUR_UTILE_CM = 17.0   # A4 (21 cm) moins deux marges de 2 cm
MOIS = ("janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
        "septembre", "octobre", "novembre", "décembre")
# Caractères interdits en XML 1.0 (contrôles hors tab/LF/CR) et demi-codets isolés
_INVALIDE_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")


# Éléments de w:pPr qui doivent suivre w:shd (ordre du schéma), et ceux qui suivent w:pBdr
_APRES_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
               "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
               "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
               "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
               "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
               "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
_APRES_SHD = _APRES_PBDR[1:]


# --- Texte ------------------------------------------------------------------
def _propre(valeur) -> str:
    """Texte sans caractère invalide en XML (python-docx échappe < > & lui-même)."""
    return _INVALIDE_XML.sub("", str(valeur))


def _txt(valeur) -> str:
    return _propre(tronquer(valeur))


def date_en_toutes_lettres(genere_le: str) -> str:
    """« 06/10/2026 à 09:05 UTC » -> « 6 octobre 2026 » (sans dépendre de la locale)."""
    m = re.match(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})", genere_le or "")
    if not m or not 1 <= int(m.group(2)) <= 12:
        return genere_le or ""
    return f"{int(m.group(1))} {MOIS[int(m.group(2)) - 1]} {m.group(3)}"


# --- Briques de mise en forme ------------------------------------------------
def _fond(element_pr, couleur: str) -> None:
    """Ajoute un fond uni (w:shd clear + fill) à un pPr ou tcPr."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)
    if element_pr.tag == qn("w:pPr"):   # respecte l'ordre du schéma (shd avant tabs/spacing/ind/jc)
        element_pr.insert_element_before(shd, *_APRES_SHD)
    else:
        element_pr.append(shd)


def _run(p, texte, taille=None, gras=False, italique=False, couleur=None):
    r = p.add_run(_propre(texte))
    r.bold, r.italic = gras or None, italique or None
    if taille:
        r.font.size = Pt(taille)
    if couleur:
        r.font.color.rgb = RGBColor.from_string(couleur)
    return r


def _para(conteneur, texte="", taille=9.5, gras=False, italique=False, couleur=TEXTE,
          alignement=None, avant=0, apres=3, retrait=0.0, suspendu=0.0, fond=None):
    """Paragraphe simple ; fond = couleur de remplissage du paragraphe."""
    p = conteneur.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(avant), Pt(apres)
    if retrait or suspendu:
        pf.left_indent, pf.first_line_indent = Cm(retrait + suspendu), Cm(-suspendu)
    if alignement is not None:
        p.alignment = alignement
    if texte:
        _run(p, texte, taille, gras, italique, couleur)
    if fond:
        _fond(p._p.get_or_add_pPr(), fond)
    return p


def _bandeau(doc, texte, couleur, taille, avant=10, apres=6, alignement=None):
    """Paragraphe pleine largeur à fond coloré, texte blanc en gras."""
    p = _para(doc, "", avant=avant, apres=apres, fond=couleur, alignement=alignement)
    p.paragraph_format.keep_with_next = True
    _run(p, texte, taille, gras=True, couleur="FFFFFF")
    # petit remplissage intérieur : bordure de la couleur du fond
    bordures = OxmlElement("w:pBdr")
    for cote in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{cote}")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "3"), ("color", couleur)):
            b.set(qn(f"w:{k}"), v)
        bordures.append(b)
    p._p.get_or_add_pPr().insert_element_before(bordures, *_APRES_PBDR)
    return p


def _section(doc, texte):
    return _bandeau(doc, texte, MARINE, 11)


def _champ(p, code: str, taille=7.5) -> None:
    """Champ Word (PAGE, NUMPAGES) : begin / instrText / separate / résultat / end."""
    def elt(type_):
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        sz, c = OxmlElement("w:sz"), OxmlElement("w:color")
        sz.set(qn("w:val"), str(int(taille * 2)))
        c.set(qn("w:val"), GRIS)
        rpr.extend([c, sz])
        r.append(rpr)
        return r

    for etat in ("begin", "instr", "separate", "texte", "end"):
        r = elt(etat)
        if etat == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {code} "
            r.append(it)
        elif etat == "texte":
            t = OxmlElement("w:t")
            t.text = "1"
            r.append(t)
        else:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), etat)
            r.append(f)
        p._p.append(r)


def _declarer_repli(doc) -> None:
    """Déclare POLICE dans la table des polices avec POLICE_REPLI comme nom
    alternatif : Word l'utilise si Poppins n'est pas installée sur le poste."""
    from lxml import etree
    for rel in doc.part.rels.values():
        if not rel.reltype.endswith("/fontTable"):
            continue
        part = rel.target_part
        racine = etree.fromstring(part.blob)
        if any(f.get(qn("w:name")) == POLICE for f in racine.findall(qn("w:font"))):
            return
        police = etree.SubElement(racine, qn("w:font"))
        police.set(qn("w:name"), POLICE)
        for balise, valeur in (("w:altName", POLICE_REPLI), ("w:family", "swiss"),
                               ("w:pitch", "variable")):
            etree.SubElement(police, qn(balise)).set(qn("w:val"), valeur)
        part._blob = etree.tostring(racine, xml_declaration=True, encoding="UTF-8",
                                    standalone=True)
        return


def _pied(section, genere_le: str) -> None:
    """Pied des pages courantes ; celui de la page de garde reste vide."""
    section.different_first_page_header_footer = True
    section.first_page_footer.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.paragraph_format.tab_stops.add_tab_stop(Cm(LARGEUR_UTILE_CM), WD_TAB_ALIGNMENT.RIGHT)
    bord = OxmlElement("w:pBdr")
    haut = OxmlElement("w:top")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "4"), ("color", TURQUOISE)):
        haut.set(qn(f"w:{k}"), v)
    bord.append(haut)
    p._p.get_or_add_pPr().insert_element_before(bord, *_APRES_PBDR)
    _run(p, f"Document généré automatiquement le {genere_le} – ForCERT", 7.5, couleur=GRIS)
    _run(p, "\tPage ", 7.5, couleur=GRIS)
    _champ(p, "PAGE")
    _run(p, " / ", 7.5, couleur=GRIS)
    _champ(p, "NUMPAGES")


# --- Tableaux ---------------------------------------------------------------
def _cellule(cell, texte, largeur_cm, fond=None, entete=False, gras=False, couleur=TEXTE):
    cell.width = Cm(largeur_cm)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    _run(p, _txt(texte) if not entete else texte, 8.5 if entete else 8, entete or gras,
         couleur="FFFFFF" if entete else couleur)
    if fond:
        _fond(cell._tc.get_or_add_tcPr(), fond)


def _tableau(doc, entetes, lignes, largeurs_cm, colonnes_gras=(), couleurs=None, alertes=()):
    """Tableau : en-tête marine répété, lignes alternées, largeurs fixes.

    largeurs_cm : None pour la colonne qui prend le reste de la largeur utile.
    couleurs : {colonne: fonction(valeur) -> couleur du texte}.
    """
    reste = LARGEUR_UTILE_CM - sum(l for l in largeurs_cm if l is not None)
    largeurs = [reste if l is None else l for l in largeurs_cm]
    t = doc.add_table(rows=1, cols=len(entetes))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, largeurs):
        col.width = Cm(w)
    couleurs = couleurs or {}
    for i, e in enumerate(entetes):
        _cellule(t.rows[0].cells[i], e, largeurs[i], MARINE, entete=True)
    trpr = t.rows[0]._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:tblHeader"))   # en-tête répété sur chaque page
    for n, ligne in enumerate(lignes):
        row = t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        fond = FOND_ALERTE if n in alertes else (LIGNE_CLAIRE if n % 2 else "FFFFFF")
        for j, v in enumerate(ligne):
            couleur = couleurs[j](v) if j in couleurs else TEXTE
            _cellule(row.cells[j], v, largeurs[j], fond, gras=j in colonnes_gras,
                     couleur=couleur)
    return t


# --- Étapes -----------------------------------------------------------------
def _ligne_etape(doc, ligne: str, condense: bool) -> None:
    """Ajoute le paragraphe d'une ligne d'explication (Recherche, Résultat, conclusion, lignes [!])."""
    retrait = len(ligne) - len(ligne.lstrip(" "))
    brut = ligne.strip()
    if not brut:
        return
    if retrait == 0:
        if condense:
            brut = _DECOMPTE_IOC.sub(r"\1.", brut)
        for prefixe, libelle in _LIBELLES:
            if brut.startswith(prefixe):
                note = prefixe == "Note :"
                p = _para(doc, apres=3, italique=note, couleur=GRIS if note else TEXTE,
                          taille=8.5 if note else 9.5)
                _run(p, libelle + " ", 8.5 if note else 9.5, gras=True,
                     couleur=GRIS if note else TURQUOISE)
                _run(p, _txt(brut[len(prefixe):].strip()), 8.5 if note else 9.5,
                     italique=note, couleur=GRIS if note else TEXTE)
                return
        _para(doc, _txt(brut))
        return
    if condense and (retrait >= 6 or brut.startswith("Note : la réputation")):
        return   # détail VirusTotal repris dans la section « Détail VirusTotal »
    if condense and (brut.startswith("- ") or _EN_TETE_PRIORITE.match(brut)):
        return   # détail repris dans le tableau qui suit l'étape
    if _EN_TETE_PRIORITE.match(brut):
        _para(doc, _txt(brut), 9, gras=True, couleur=MARINE, avant=3, apres=2, retrait=0.4)
        return
    if brut.startswith("[!]"):
        p = _para(doc, apres=2, retrait=0.5, suspendu=0.7, fond=FOND_ALERTE)
        _run(p, "[!]\t", 9, gras=True, couleur=ALERTE)
        _run(p, _txt(brut[3:].strip()), 9)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.2))
        return
    puce = brut.startswith("- ")
    texte = brut[2:] if puce else brut
    gauche = 0.4 if retrait < 4 else 1.0
    p = _para(doc, apres=2, retrait=gauche + (0.0 if puce else 0.5), suspendu=0.5 if puce else 0,
              taille=9 if retrait < 4 else 8.5)
    if puce:
        _run(p, "•\t", 9)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(gauche + 0.5))
    _run(p, _txt(texte), 9 if retrait < 4 else 8.5)


def _etape(doc, etape: dict, numero: int, condense=False, renvoi="") -> None:
    titre = re.sub(r"^ÉTAPE\s+\d+\s*:\s*", "", etape["titre"])
    _section(doc, f"Étape {numero} : {_txt(titre)}")
    for ligne in etape["texte"].splitlines():
        _ligne_etape(doc, ligne, condense)
    if renvoi:
        _para(doc, renvoi, 8.5, italique=True, couleur=GRIS)


# --- Document ---------------------------------------------------------------
def _page_de_garde(doc, d: dict, logo) -> None:
    if logo and os.path.isfile(logo):
        p = _para(doc, alignement=WD_ALIGN_PARAGRAPH.CENTER, apres=0)
        try:
            p.add_run().add_picture(logo, width=Cm(7))
        except Exception:   # image illisible : page de garde sans logo
            p._p.clear_content()
            print(f"Avertissement : logo illisible ({logo}), page de garde sans logo.",
                  file=sys.stderr)
    _para(doc, _txt(d["titre"]), 28, gras=True, couleur=MARINE, avant=70, apres=14,
          alignement=WD_ALIGN_PARAGRAPH.CENTER)
    _bandeau(doc, "Investigation automatisée des journaux Windows", TURQUOISE, 14,
             avant=0, apres=36, alignement=WD_ALIGN_PARAGRAPH.CENTER)
    centre = WD_ALIGN_PARAGRAPH.CENTER
    _para(doc, date_en_toutes_lettres(d["genere_le"]), 16, gras=True, couleur=MARINE,
          apres=14, alignement=centre)
    periode = d["periode"]
    _para(doc, "Période analysée : " + ("du " if periode[:1].isdigit() else "") + _txt(periode),
          10.5, couleur=GRIS, alignement=centre, apres=4)
    if d.get("fichier"):
        n = d.get("nb_evenements", 0)
        _para(doc, f"Source : {_txt(d['fichier'])} ({n} événement{'s' if n > 1 else ''}) "
                   f"· heures en UTC", 10.5, couleur=GRIS, alignement=centre, apres=36)
    g = d["gravite"]
    fond = COULEURS_GRAVITE.get(g, GRIS)
    texte = MARINE if g == "MODÉRÉE" else "FFFFFF"
    p = _para(doc, alignement=centre, fond=fond, avant=0, apres=0)
    _run(p, f"GRAVITÉ : {_propre(g)}", 18, gras=True, couleur=texte)
    _para(doc, "Document généré automatiquement – ForCERT", 9, italique=True, couleur=GRIS,
          avant=150, alignement=centre)
    doc.add_page_break()


def generer_docx(chemin: str, donnees: dict, logo: str | None = None) -> None:
    """Écrit le rapport d'incident Word (A4) à partir de construire_donnees_rapport."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    sec.footer_distance = Cm(0.9)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = POLICE, Pt(9.5)
    rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), POLICE)
    _declarer_repli(doc)
    for zoom in doc.settings.element.findall(qn("w:zoom")):   # gabarit python-docx : percent manquant
        zoom.set(qn("w:percent"), "100")
    cp = doc.core_properties
    cp.title, cp.author = _propre(donnees["titre"]), "ForCERT"
    cp.subject = "Rapport d'incident généré automatiquement"

    _pied(sec, _propre(donnees["genere_le"]))
    _page_de_garde(doc, donnees, logo)

    _section(doc, "En bref")
    for par in donnees["synthese"]:
        _para(doc, _txt(par), 11, alignement=WD_ALIGN_PARAGRAPH.JUSTIFY, apres=6)

    _section(doc, "Actions à lancer immédiatement")
    actions = donnees["actions_prioritaires"]
    if actions:
        for i, a in enumerate(actions, 1):
            p = _para(doc, apres=4, retrait=0.0, suspendu=0.7)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.7))
            _run(p, f"{i}.\t", 10, gras=True, couleur=TURQUOISE)
            _run(p, _txt(a), 10)
    else:
        _para(doc, "Aucune action immédiate n'est nécessaire.")
    _para(doc, "Les pages suivantes détaillent le raisonnement suivi, étape par étape, la "
               "chronologie de l'attaque, le plan d'action complet et, en annexe, les "
               "indicateurs techniques à transmettre à l'équipe sécurité.",
          8.5, italique=True, couleur=GRIS, avant=6)
    doc.add_page_break()

    etapes = donnees["etapes"]
    for i, etape in enumerate(etapes[:3], 1):
        _etape(doc, etape, i)

    _section(doc, "Chronologie de l'attaque")
    if donnees["chronologie"]:
        lignes = [[h, m, c, d[4:] if d.startswith("[!] ") else d]
                  for h, m, c, d in donnees["chronologie"]]
        suspects = {i for i, brut in enumerate(donnees["chronologie"])
                    if brut[3].startswith("[!] ")}
        large = any(len(l[0]) > 5 for l in lignes)
        _tableau(doc, ["Heure", "Machine", "Compte", "Ce qui s'est passé"], lignes,
                 [3.4 if large else 1.8, 3.0, 2.8, None], colonnes_gras={0},
                 alertes=suspects)
        _para(doc, "Les actions suspectes sont surlignées en rouge.", 8.5, italique=True,
              couleur=GRIS, avant=3)
    else:
        _para(doc, "Aucune activité d'attaquant à retracer.")

    if len(etapes) > 3:
        _etape(doc, etapes[3], 4, True, "Le détail des actions figure dans le tableau ci-dessous.")
    _section(doc, "Plan d'action")
    if donnees["plan"]:
        _tableau(doc, ["Fait observé", "Action à mener", "Priorité"], donnees["plan"],
                 [5.8, None, 2.6], colonnes_gras={2},
                 couleurs={2: lambda v: COULEURS_PRIORITE.get(v, MARINE)})
    else:
        _para(doc, "Aucune action à mener.")

    if len(etapes) > 4:
        _etape(doc, etapes[4], 5, True, "La liste des indicateurs figure dans l'annexe ci-dessous.")
    _section(doc, "Annexe : indicateurs de compromission")
    if donnees["iocs"]:
        _tableau(doc, ["Indicateur", "Type", "Vérification VirusTotal"], donnees["iocs"],
                 [8.2, 3.2, None])
    else:
        _para(doc, "Aucun indicateur extrait.")
    if donnees.get("details_vt"):
        _section(doc, "Détail VirusTotal")
        for d in donnees["details_vt"]:
            p = _para(doc, "", avant=4, apres=1)
            p.paragraph_format.keep_with_next = True
            _run(p, _txt(f"{d['valeur']} — {d['libelle']}"), 9.5, gras=True, couleur=MARINE)
            for ligne in d["lignes"]:
                _para(doc, _txt(ligne), 8.5, retrait=0.5, apres=1)
    _para(doc, PRUDENCE, 8.5, italique=True, couleur=GRIS, avant=8, fond="FFF8E6")

    doc.save(chemin)
