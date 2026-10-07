import base64
import contextlib
import io
import importlib.util
import os
import shutil
import tempfile
import unittest
from datetime import datetime, timezone

from investigation import construire_donnees_rapport, investiguer
from tests.jeu_synthetique import IP_SPRAY_A, construire_jeu, ecrire_jeu

DOCX = importlib.util.find_spec("docx") is not None
GENERE_LE = datetime(2026, 10, 6, 9, 5, tzinfo=timezone.utc)
LOGO = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "Logo.png")


@unittest.skipUnless(DOCX, "python-docx absent")
class TestRapportWord(unittest.TestCase):
    def setUp(self):
        self.dossier = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.dossier, True)

    def generer(self, evenements=None, logo=None):
        from docx import Document
        from rapport_word import generer_docx
        chemin_jeu = ecrire_jeu(evenements)
        self.addCleanup(os.remove, chemin_jeu)
        donnees = construire_donnees_rapport(investiguer(chemin_jeu), GENERE_LE)
        sortie = os.path.join(self.dossier, "r.docx")
        generer_docx(sortie, donnees, logo)
        return Document(sortie)

    @staticmethod
    def texte(doc):
        morceaux = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for ligne in t.rows:
                morceaux += [c.text for c in ligne.cells]
        return "\n".join(morceaux)

    def test_charte_couleurs_et_police(self):
        # Couleurs et police relevées sur la charte Formind ; repli Arial déclaré.
        from docx.oxml.ns import qn
        from lxml import etree
        doc = self.generer()
        corps = etree.tostring(doc.element).decode()
        self.assertIn('w:fill="002236"', corps)
        self.assertIn('w:fill="0089A4"', corps)
        self.assertNotIn("1B2A4A", corps)
        self.assertEqual(doc.styles["Normal"].font.name, "Poppins")
        table = next(r.target_part for r in doc.part.rels.values()
                     if r.reltype.endswith("/fontTable"))
        racine = etree.fromstring(table.blob)
        poppins = [f for f in racine.findall(qn("w:font")) if f.get(qn("w:name")) == "Poppins"]
        self.assertEqual(len(poppins), 1)
        self.assertEqual(poppins[0].find(qn("w:altName")).get(qn("w:val")), "Arial")

    def test_contenu_principal(self):
        t = self.texte(self.generer())
        for attendu in ("Rapport d'incident de sécurité",
                        "Investigation automatisée des journaux Windows",
                        "6 octobre 2026", "GRAVITÉ : CRITIQUE", "En bref",
                        "Chronologie de l'attaque", "Plan d'action",
                        "Annexe : indicateurs de compromission", "ForCERT",
                        "Rappel de prudence", IP_SPRAY_A, "adm_tmp"):
            self.assertIn(attendu, t)

    def test_logo(self):
        self.assertEqual(len(self.generer(logo=LOGO).inline_shapes), 1)
        self.assertEqual(len(self.generer(logo=None).inline_shapes), 0)
        self.assertEqual(len(self.generer(logo="/inexistant/x.png").inline_shapes), 0)

    def test_tableaux_et_entete_repetee(self):
        doc = self.generer()
        self.assertEqual(len(doc.tables), 3)
        for t in doc.tables:
            self.assertEqual(len(t._tbl.xpath("./w:tr[1]/w:trPr/w:tblHeader")), 1)

    def test_page_de_garde_sans_pied_numerote(self):
        doc = self.generer()
        sec = doc.sections[0]
        self.assertTrue(sec.different_first_page_header_footer)
        self.assertEqual(sec.first_page_footer.paragraphs[0].text, "")
        xml = sec.footer._element.xml
        self.assertIn("NUMPAGES", xml)
        self.assertIn("PAGE", xml)

    def test_texte_a_echapper(self):
        evts = construire_jeu()
        for e in evts:
            if e["event_id"] == 4698:
                e["details"]["task_name"] = "\\A<b>&C"
        self.assertIn("\\A<b>&C", self.texte(self.generer(evts)))

    def test_commande_tres_longue(self):
        evts = construire_jeu()
        charge = base64.b64encode(("Write-Host 'x' <&>\\ " * 230).encode("utf-16-le")).decode()
        for e in evts:
            if e["event_id"] == 4688 and e["account"] == "adm_tmp":
                e["details"]["command_line"] = "powershell.exe -enc " + charge
        self.assertGreater(len(charge), 9000)
        self.assertIn("tronqué", self.texte(self.generer(evts)))

    def test_ordre_xml_et_logo_corrompu(self):
        doc = self.generer()
        for t in doc.tables:
            self.assertEqual(len(t._tbl.xpath("./w:tblPr/w:tblLayout")), 1)
        bandeau = next(p for p in doc.paragraphs if p.text == "En bref")
        noms = [c.tag.split("}")[1] for c in bandeau._p.pPr]
        self.assertEqual(noms, sorted(noms, key=["keepNext", "pBdr", "shd", "spacing"].index))
        corrompu = os.path.join(self.dossier, "mauvais.png")
        with open(corrompu, "wb") as f:
            f.write(b"pas une image")
        with contextlib.redirect_stderr(io.StringIO()) as err:
            doc = self.generer(logo=corrompu)
        self.assertEqual(len(doc.inline_shapes), 0)
        self.assertIn("logo", err.getvalue())

    def test_caracteres_de_controle(self):
        from rapport_word import _propre
        self.assertEqual(_propre("a\x00b\x1bc\ud800d"), "abcd")

    def test_jeu_vide(self):
        self.assertIn("En bref", self.texte(self.generer([])))


if __name__ == "__main__":
    unittest.main()


class TestDetailsVTWord(unittest.TestCase):
    def test_details_vt_rendus(self):
        from docx import Document
        from rapport_word import generer_docx
        d = construire_donnees_rapport(investiguer(ecrire_jeu()), GENERE_LE)
        d["details_vt"] = [{"valeur": "203.0.113.47", "libelle": "IP d'attaque",
                            "lignes": ["Pays : États-Unis (US)", "Fiche VirusTotal : https://x"]}]
        chemin = os.path.join(tempfile.mkdtemp(), "r.docx")
        generer_docx(chemin, d)
        t = "\n".join(p.text for p in Document(chemin).paragraphs)
        for a in ("Détail VirusTotal", "203.0.113.47 — IP d'attaque", "États-Unis (US)"):
            self.assertIn(a, t)
        d["details_vt"] = []
        generer_docx(chemin, d)
        self.assertNotIn("Détail VirusTotal",
                         "\n".join(p.text for p in Document(chemin).paragraphs))


class TestEtape5CondenseeWord(unittest.TestCase):
    def test_details_vt_seulement_dans_leur_section(self):
        from docx import Document
        from rapport_word import generer_docx
        d = construire_donnees_rapport(investiguer(ecrire_jeu()), GENERE_LE)
        d["details_vt"] = [{"valeur": "203.0.113.47", "libelle": "IP d'attaque",
                            "lignes": ["Pays : États-Unis (US)"]}]
        for e in d["etapes"]:
            if e["titre"].startswith("ÉTAPE 5"):
                e["texte"] += ("\n      Pays : États-Unis (US)\n      Analyses : 1 malveillant"
                               "\n  Note : la réputation est le score attribué par la communauté")
        chemin = os.path.join(tempfile.mkdtemp(), "r.docx")
        generer_docx(chemin, d)
        t = "\n".join(p.text for p in Document(chemin).paragraphs)
        self.assertEqual(t.count("Pays :"), 1)
        self.assertNotIn("Analyses :", t)
        self.assertNotIn("la réputation est le score", t)
